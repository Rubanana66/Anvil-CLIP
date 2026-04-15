from __future__ import annotations

import argparse
import json
import logging
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.datavalidation import DataValidation

import icon_matcher
import parser as stockpile_parser
from review_matches import add_image, load_item_names, read_number_crop


LOGGER = logging.getLogger("anvil_batch_review")
IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}


@dataclass(frozen=True)
class BatchSlotRow:
    screenshot_name: str
    parser_output_dir: Path
    slot_index: int
    row: int
    col: int
    slot_path: Path
    icon_path: Path
    number_path: Path
    detected_amount: str
    detected_item: str
    icon_confidence: float
    review_needed: bool
    top_candidates: str

    def to_dict(self) -> dict[str, Any]:
        return {
            "screenshot_name": self.screenshot_name,
            "parser_output_dir": str(self.parser_output_dir),
            "slot_index": self.slot_index,
            "row": self.row,
            "col": self.col,
            "slot_path": str(self.slot_path),
            "icon_path": str(self.icon_path),
            "number_path": str(self.number_path),
            "detected_amount": self.detected_amount,
            "detected_item": self.detected_item,
            "icon_confidence": self.icon_confidence,
            "review_needed": self.review_needed,
            "top_candidates": self.top_candidates,
        }


@dataclass(frozen=True)
class BatchError:
    screenshot_name: str
    error: str

    def to_dict(self) -> dict[str, str]:
        return {
            "screenshot_name": self.screenshot_name,
            "error": self.error,
        }


def discover_screenshots(screenshots_dir: Path) -> list[Path]:
    if screenshots_dir.is_file():
        return [screenshots_dir]

    if not screenshots_dir.exists():
        raise FileNotFoundError(f"Screenshots folder does not exist: {screenshots_dir}")

    return sorted(
        (
            path
            for path in screenshots_dir.iterdir()
            if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS
        ),
        key=lambda path: path.name.casefold(),
    )


def run_batch(
    screenshots_dir: Path,
    icon_dir: Path,
    output_dir: Path,
    model_name: str,
    pretrained: str,
    device: str | None,
    top_k: int,
    review_threshold: float,
    margin_threshold: float,
) -> tuple[Path, list[BatchSlotRow], list[BatchError]]:
    screenshots = discover_screenshots(screenshots_dir)
    if not screenshots:
        raise FileNotFoundError(f"No screenshots found in {screenshots_dir}")

    timestamp = datetime.now().strftime("%Y%m%d_%H%M%S")
    batch_dir = output_dir / f"combined_review_{timestamp}"
    parser_base_dir = batch_dir / "parsed"
    batch_dir.mkdir(parents=True, exist_ok=True)

    LOGGER.info("Found %d screenshot(s).", len(screenshots))
    LOGGER.info("Batch output folder: %s", batch_dir)

    parsed_results: list[tuple[Path, stockpile_parser.ParseResult]] = []
    errors: list[BatchError] = []

    for screenshot_path in screenshots:
        try:
            LOGGER.info("Parsing %s", screenshot_path.name)
            result = stockpile_parser.parse_screenshot(screenshot_path, parser_base_dir)
            parsed_results.append((screenshot_path, result))
            LOGGER.info("Detected %d slot(s) in %s", result.slot_count, screenshot_path.name)
        except Exception as exc:
            LOGGER.exception("Failed to parse %s", screenshot_path)
            errors.append(BatchError(screenshot_path.name, str(exc)))

    icon_paths = [
        slot.icon_path.resolve()
        for _, result in parsed_results
        for slot in result.slots
    ]

    matches_by_path: dict[Path, icon_matcher.IconMatch] = {}
    if icon_paths:
        matcher = icon_matcher.IconMatcher(
            icon_dir=icon_dir,
            model_name=model_name,
            pretrained=pretrained,
            device=device,
            review_threshold=review_threshold,
            margin_threshold=margin_threshold,
        )
        matches = matcher.match_many(icon_paths, top_k=top_k)
        matches_by_path = {match.query_path.resolve(): match for match in matches}
        LOGGER.info("Matched %d icon crop(s).", len(matches))

        by_parser_dir: dict[Path, list[icon_matcher.IconMatch]] = {}
        for match in matches:
            parser_output_dir = match.query_path.resolve().parent.parent
            by_parser_dir.setdefault(parser_output_dir, []).append(match)

        for parser_output_dir, parser_matches in by_parser_dir.items():
            matches_dir = parser_output_dir / "icon_matches"
            icon_matcher.write_json(parser_matches, matches_dir / "icon_matches.json")
            icon_matcher.write_csv(parser_matches, matches_dir / "icon_matches.csv")

    rows: list[BatchSlotRow] = []
    for screenshot_path, result in parsed_results:
        parser_output_dir = _parser_output_dir(result)
        for slot in result.slots:
            icon_path = slot.icon_path.resolve()
            slot_path = slot.slot_path.resolve()
            number_path = slot.number_path.resolve()
            match = matches_by_path.get(icon_path)

            detected_amount = read_number_crop(number_path)
            detected_item = match.best_item if match else ""
            icon_confidence = match.similarity if match else 0.0
            matcher_review = match.review_needed if match else True
            amount_review = detected_amount == ""
            top_candidates = _format_candidates(match)

            rows.append(
                BatchSlotRow(
                    screenshot_name=screenshot_path.name,
                    parser_output_dir=parser_output_dir,
                    slot_index=slot.slot_index,
                    row=slot.row,
                    col=slot.col,
                    slot_path=slot_path,
                    icon_path=icon_path,
                    number_path=number_path,
                    detected_amount=detected_amount,
                    detected_item=detected_item,
                    icon_confidence=icon_confidence,
                    review_needed=matcher_review or amount_review,
                    top_candidates=top_candidates,
                )
            )

    workbook_path = batch_dir / "combined_stockpile_review.xlsx"
    write_combined_workbook(workbook_path, rows, errors, icon_dir)
    write_summary(batch_dir / "batch_summary.json", screenshots, rows, errors, workbook_path)
    return workbook_path, rows, errors


def write_combined_workbook(
    output_path: Path,
    rows: list[BatchSlotRow],
    errors: list[BatchError],
    icon_dir: Path,
) -> None:
    item_names = load_item_names(icon_dir)

    workbook = Workbook()
    review = workbook.active
    review.title = "Review"

    headers = [
        "Screenshot",
        "Slot",
        "Row",
        "Col",
        "Slot Crop",
        "Icon Crop",
        "Detected Amount",
        "Detected Item",
        "Icon Confidence",
        "Review Needed",
        "Top Candidates",
        "Correct Item",
        "Correct Amount",
        "Notes",
        "Parser Output",
        "Icon Path",
        "Slot Path",
        "Number Path",
    ]
    review.append(headers)
    _style_header(review)

    for row_number, row in enumerate(rows, start=2):
        review.append(
            [
                row.screenshot_name,
                row.slot_index,
                row.row,
                row.col,
                "",
                "",
                row.detected_amount,
                row.detected_item,
                row.icon_confidence,
                row.review_needed,
                row.top_candidates,
                row.detected_item,
                row.detected_amount,
                "",
                str(row.parser_output_dir),
                str(row.icon_path),
                str(row.slot_path),
                str(row.number_path),
            ]
        )
        review.row_dimensions[row_number].height = 74
        add_image(review, row.slot_path, f"E{row_number}", width=72, height=72)
        add_image(review, row.icon_path, f"F{row_number}", width=72, height=72)

        review.cell(row=row_number, column=9).number_format = "0.0000"
        for column in range(1, len(headers) + 1):
            review.cell(row=row_number, column=column).alignment = Alignment(vertical="top", wrap_text=True)

    items_sheet = workbook.create_sheet("Item Names")
    for row_number, item_name in enumerate(item_names, start=1):
        items_sheet.cell(row=row_number, column=1, value=item_name)
    items_sheet.sheet_state = "hidden"

    if item_names and rows:
        validation = DataValidation(
            type="list",
            formula1=f"'Item Names'!$A$1:$A${len(item_names)}",
            allow_blank=True,
        )
        review.add_data_validation(validation)
        validation.add(f"L2:L{len(rows) + 1}")

    if errors:
        error_sheet = workbook.create_sheet("Errors")
        error_sheet.append(["Screenshot", "Error"])
        _style_header(error_sheet)
        for error in errors:
            error_sheet.append([error.screenshot_name, error.error])
        error_sheet.column_dimensions["A"].width = 48
        error_sheet.column_dimensions["B"].width = 100
        for row in error_sheet.iter_rows():
            for cell in row:
                cell.alignment = Alignment(vertical="top", wrap_text=True)

    widths = {
        "A": 42,
        "B": 8,
        "C": 8,
        "D": 8,
        "E": 14,
        "F": 14,
        "G": 16,
        "H": 28,
        "I": 16,
        "J": 15,
        "K": 72,
        "L": 28,
        "M": 16,
        "N": 32,
        "O": 58,
        "P": 58,
        "Q": 58,
        "R": 58,
    }
    for column, width in widths.items():
        review.column_dimensions[column].width = width

    review.freeze_panes = "A2"
    review.auto_filter.ref = review.dimensions
    output_path.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output_path)


def write_summary(
    output_path: Path,
    screenshots: list[Path],
    rows: list[BatchSlotRow],
    errors: list[BatchError],
    workbook_path: Path,
) -> None:
    summary = {
        "created_at": datetime.now().isoformat(timespec="seconds"),
        "workbook_path": str(workbook_path),
        "screenshots": [str(path) for path in screenshots],
        "slot_count": len(rows),
        "error_count": len(errors),
        "rows": [row.to_dict() for row in rows],
        "errors": [error.to_dict() for error in errors],
    }
    output_path.write_text(json.dumps(summary, indent=2), encoding="utf-8")


def _style_header(sheet: Any) -> None:
    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _format_candidates(match: icon_matcher.IconMatch | None) -> str:
    if not match:
        return ""
    return " | ".join(
        f"{candidate.item_name}:{candidate.similarity:.4f}"
        for candidate in match.candidates
    )


def _parser_output_dir(result: stockpile_parser.ParseResult) -> Path:
    json_path = result.debug_paths.get("json_summary")
    if json_path:
        return json_path.resolve().parent
    if result.slots:
        return result.slots[0].slot_path.resolve().parent.parent
    original_path = result.debug_paths.get("original")
    if original_path:
        return original_path.resolve().parent
    return Path.cwd()


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Parse all stockpile screenshots and build one Excel review workbook.")
    parser.add_argument(
        "screenshots_dir",
        nargs="?",
        type=Path,
        default=Path("screenshots"),
        help="Folder containing stockpile screenshots. Default: screenshots",
    )
    parser.add_argument(
        "--icon-dir",
        type=Path,
        default=Path("itemlist"),
        help="Folder containing clean item PNGs. Default: itemlist",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path("output_debug"),
        help="Base folder for debug output and the combined workbook. Default: output_debug",
    )
    parser.add_argument(
        "--model",
        default=icon_matcher.MODEL_NAME,
        help=f"OpenCLIP model name. Default: {icon_matcher.MODEL_NAME}",
    )
    parser.add_argument(
        "--pretrained",
        default=icon_matcher.PRETRAINED,
        help=f"OpenCLIP pretrained weights. Default: {icon_matcher.PRETRAINED}",
    )
    parser.add_argument(
        "--device",
        choices=("cpu", "cuda"),
        help="Torch device. Default: cuda if available, otherwise cpu.",
    )
    parser.add_argument(
        "--top-k",
        type=int,
        default=5,
        help="Number of candidates to keep per icon. Default: 5",
    )
    parser.add_argument(
        "--review-threshold",
        type=float,
        default=0.24,
        help="Flag review if best similarity is below this. Default: 0.24",
    )
    parser.add_argument(
        "--margin-threshold",
        type=float,
        default=0.025,
        help="Flag review if top match barely beats second match. Default: 0.025",
    )
    parser.add_argument(
        "--log-level",
        default="INFO",
        choices=("DEBUG", "INFO", "WARNING", "ERROR"),
        help="Logging level. Default: INFO",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    logging.basicConfig(level=getattr(logging, args.log_level), format="%(levelname)s: %(message)s")

    workbook_path, rows, errors = run_batch(
        screenshots_dir=args.screenshots_dir,
        icon_dir=args.icon_dir,
        output_dir=args.output_dir,
        model_name=args.model,
        pretrained=args.pretrained,
        device=args.device,
        top_k=args.top_k,
        review_threshold=args.review_threshold,
        margin_threshold=args.margin_threshold,
    )

    print()
    print("Anvil combined stockpile review")
    print("--------------------------------")
    print(f"Excel: {workbook_path}")
    print(f"Rows: {len(rows)}")
    print(f"Errors: {len(errors)}")


if __name__ == "__main__":
    main()
