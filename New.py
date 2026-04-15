    from __future__ import annotations

import argparse
import csv
import re
import shutil
import sys
from dataclasses import dataclass
from datetime import date
from pathlib import Path

try:
    import cv2
    import imagehash
    import numpy as np
    import pytesseract
    from openpyxl import Workbook, load_workbook
    from PIL import Image, ImageDraw
except ImportError as exc:
    missing = exc.name or "a required package"
    print(
        f"Missing dependency: {missing}\n\n"
        "Install/update the Python packages with:\n"
        "  py -m pip install -r requirements.txt\n\n"
        "You also need the Tesseract OCR app installed for number reading:\n"
        "  https://github.com/UB-Mannheim/tesseract/wiki\n",
        file=sys.stderr,
    )
    raise SystemExit(1)


DEFAULT_IMAGES_DIR = Path("anvil_item_images")
DEFAULT_ITEMS_WORKBOOK = Path("anvil_items.xlsx")
DEFAULT_OUTPUT = Path("family_keep_inventory.xlsx")
DEFAULT_ANNOTATIONS_DIR = Path("family_keep_review")
DEFAULT_SCREENSHOTS_DIR = Path("screenshots")
WINDOWS_TESSERACT_PATHS = (
    Path(r"C:\Program Files\Tesseract-OCR\tesseract.exe"),
    Path(r"C:\Program Files (x86)\Tesseract-OCR\tesseract.exe"),
)
MIN_SLOT_SIZE = 45
MAX_SLOT_SIZE = 180


@dataclass(frozen=True)
class Slot:
    x: int
    y: int
    w: int
    h: int


@dataclass
class ItemReference:
    name: str
    image_path: Path
    phash: imagehash.ImageHash
    dhash: imagehash.ImageHash
    colorhash: imagehash.ImageHash
    template_rgba: np.ndarray


@dataclass
class Detection:
    slot: Slot
    item_name: str
    amount: int
    confidence: float
    image_path: Path
    screenshot: Path
    candidates: str = ""


def item_name_from_png(path: Path) -> str:
    # Created by Anvil.py as Item_Name_ab12cd34.png.
    stem = re.sub(r"_[0-9a-fA-F]{8}$", "", path.stem)
    return stem.replace("_", " ")


def trim_alpha(image: Image.Image) -> Image.Image:
    rgba = image.convert("RGBA")
    alpha_box = rgba.getchannel("A").getbbox()
    if alpha_box:
        return rgba.crop(alpha_box)
    return rgba


def template_rgba(image: Image.Image) -> np.ndarray:
    rgba = trim_alpha(image)
    return np.asarray(rgba.convert("RGBA"))


def normalise_for_hash(image: Image.Image, canvas_size: int = 96) -> Image.Image:
    image = trim_alpha(image)
    image.thumbnail((canvas_size - 12, canvas_size - 12), Image.Resampling.LANCZOS)

    canvas = Image.new("RGBA", (canvas_size, canvas_size), (82, 80, 70, 255))
    x = (canvas_size - image.width) // 2
    y = (canvas_size - image.height) // 2
    canvas.alpha_composite(image.convert("RGBA"), (x, y))
    return canvas.convert("RGB")


def crop_slot_icon(slot_image: Image.Image) -> Image.Image:
    image = slot_image.convert("RGB")
    width, height = image.size

    # Hide the bottom-right count text before matching the icon.
    draw = ImageDraw.Draw(image)
    patch = image.crop((0, 0, max(4, width // 6), max(4, height // 6)))
    background = tuple(int(v) for v in np.asarray(patch).reshape(-1, 3).mean(axis=0))
    draw.rectangle((int(width * 0.52), int(height * 0.56), width, height), fill=background)

    return normalise_for_hash(image)


def image_fingerprint(image: Image.Image) -> tuple[imagehash.ImageHash, imagehash.ImageHash, imagehash.ImageHash]:
    return (
        imagehash.phash(image, hash_size=12),
        imagehash.dhash(image, hash_size=12),
        imagehash.colorhash(image),
    )


def configure_tesseract(tesseract_cmd: str | None = None) -> str | None:
    if tesseract_cmd:
        pytesseract.pytesseract.tesseract_cmd = tesseract_cmd
        return tesseract_cmd

    discovered = shutil.which("tesseract")
    if discovered:
        pytesseract.pytesseract.tesseract_cmd = discovered
        return discovered

    for candidate in WINDOWS_TESSERACT_PATHS:
        if candidate.exists():
            pytesseract.pytesseract.tesseract_cmd = str(candidate)
            return str(candidate)

    return None


def tesseract_help_message() -> str:
    return (
        "Tesseract OCR is needed to read item stack numbers, but it was not found.\n\n"
        "Install the Windows build from:\n"
        "  https://github.com/UB-Mannheim/tesseract/wiki\n\n"
        "Then restart the app. If it still is not found, enter this in the app sidebar:\n"
        r"  C:\Program Files\Tesseract-OCR\tesseract.exe"
    )


def hash_distance(
    left: tuple[imagehash.ImageHash, imagehash.ImageHash, imagehash.ImageHash],
    right: ItemReference,
) -> float:
    phash, dhash, colorhash = left
    return float((phash - right.phash) + (0.65 * (dhash - right.dhash)) + (4.0 * (colorhash - right.colorhash)))


def template_match_score(slot_image: Image.Image, reference: ItemReference) -> float:
    slot = np.asarray(slot_image.convert("RGB"))
    slot_gray = cv2.cvtColor(slot, cv2.COLOR_RGB2GRAY)
    template = getattr(reference, "template_rgba", None)
    if template is None:
        template = template_rgba(Image.open(reference.image_path))
        reference.template_rgba = template

    alpha = template[:, :, 3]
    if int(np.count_nonzero(alpha > 20)) < 40:
        return 0.0

    template_rgb = template[:, :, :3]
    template_gray = cv2.cvtColor(template_rgb, cv2.COLOR_RGB2GRAY)
    mask = (alpha > 20).astype(np.uint8) * 255

    slot_h, slot_w = slot_gray.shape[:2]
    original_h, original_w = template_gray.shape[:2]
    if original_h == 0 or original_w == 0:
        return 0.0

    best = 0.0
    target_sides = [
        int(min(slot_w, slot_h) * scale)
        for scale in (0.52, 0.62, 0.72, 0.82, 0.92)
    ]

    for target_side in target_sides:
        scale = target_side / max(original_w, original_h)
        resized_w = max(6, int(original_w * scale))
        resized_h = max(6, int(original_h * scale))
        if resized_w > slot_w or resized_h > slot_h:
            continue

        resized_template = cv2.resize(template_gray, (resized_w, resized_h), interpolation=cv2.INTER_AREA)
        resized_mask = cv2.resize(mask, (resized_w, resized_h), interpolation=cv2.INTER_NEAREST)

        try:
            result = cv2.matchTemplate(
                slot_gray,
                resized_template,
                cv2.TM_CCOEFF_NORMED,
                mask=resized_mask,
            )
        except cv2.error:
            continue

        _, max_value, _, _ = cv2.minMaxLoc(result)
        if np.isfinite(max_value):
            best = max(best, float(max_value))

    return best


def load_allowed_names(item_names_csv: Path | None, item_names_workbook: Path | None) -> set[str] | None:
    allowed_names: set[str] = set()

    if item_names_csv and item_names_csv.exists():
        with item_names_csv.open(newline="", encoding="utf-8") as handle:
            reader = csv.DictReader(handle)
            for row in reader:
                name = row.get("Item Name")
                if name:
                    allowed_names.add(name)

    if item_names_workbook and item_names_workbook.exists():
        workbook = load_workbook(item_names_workbook, read_only=True, data_only=True)
        sheet = workbook.active
        headers = [str(cell.value or "") for cell in sheet[1]]
        try:
            name_column = headers.index("Item Name") + 1
        except ValueError:
            name_column = 2

        for row in sheet.iter_rows(min_row=2, values_only=True):
            if len(row) >= name_column and row[name_column - 1]:
                allowed_names.add(str(row[name_column - 1]))

    return allowed_names or None


def load_references(
    images_dir: Path,
    item_names_csv: Path | None = None,
    item_names_workbook: Path | None = DEFAULT_ITEMS_WORKBOOK,
) -> list[ItemReference]:
    if not images_dir.exists():
        raise FileNotFoundError(
            f"Could not find {images_dir}. Run Anvil.py first so the item PNGs exist."
        )

    allowed_names = load_allowed_names(item_names_csv, item_names_workbook)

    references: list[ItemReference] = []
    for image_path in sorted(images_dir.glob("*.png")):
        name = item_name_from_png(image_path)
        if allowed_names is not None and name not in allowed_names:
            continue

        try:
            raw_image = Image.open(image_path)
            image = normalise_for_hash(raw_image)
            phash, dhash, colorhash = image_fingerprint(image)
            template = template_rgba(raw_image)
        except Exception as exc:
            print(f"Skipping reference {image_path}: {exc}", file=sys.stderr)
            continue

        references.append(ItemReference(name, image_path, phash, dhash, colorhash, template))

    if not references:
        raise RuntimeError(f"No usable item PNGs found in {images_dir}.")

    return references


def find_slots(screenshot_path: Path) -> list[Slot]:
    image = cv2.imread(str(screenshot_path))
    if image is None:
        raise FileNotFoundError(f"Could not open screenshot: {screenshot_path}")

    gray = cv2.cvtColor(image, cv2.COLOR_BGR2GRAY)
    edges = cv2.Canny(gray, 50, 120)
    contours, _ = cv2.findContours(edges, cv2.RETR_LIST, cv2.CHAIN_APPROX_SIMPLE)

    slots: list[Slot] = []
    for contour in contours:
        x, y, w, h = cv2.boundingRect(contour)
        aspect = w / max(h, 1)
        area = w * h
        if not (MIN_SLOT_SIZE <= w <= MAX_SLOT_SIZE and MIN_SLOT_SIZE <= h <= MAX_SLOT_SIZE):
            continue
        if not (0.72 <= aspect <= 1.35):
            continue
        if area < 2500:
            continue
        slots.append(Slot(x, y, w, h))

    # Remove duplicate rectangles found on the same border.
    slots = sorted(slots, key=slot_score, reverse=True)
    deduped: list[Slot] = []
    for slot in slots:
        if any(slot_iou(slot, old) > 0.42 or slot_is_inside(slot, old) for old in deduped):
            continue
        deduped.append(slot)

    return sorted(deduped, key=lambda slot: (slot.y, slot.x))


def slot_score(slot: Slot) -> float:
    width_score = 1.0 - min(abs(slot.w - 102), 80) / 80
    height_score = 1.0 - min(abs(slot.h - 102), abs(slot.h - 86), 80) / 80
    aspect_score = 1.0 - min(abs((slot.w / max(slot.h, 1)) - 1.0), 0.6) / 0.6
    too_large_penalty = 1.5 if slot.w > 122 or slot.h > 122 else 0.0
    too_small_penalty = 1.0 if slot.w < 76 or slot.h < 76 else 0.0
    edge_penalty = 0.8 if slot.x < 10 else 0.0
    return (width_score * 3.0) + (height_score * 3.0) + (aspect_score * 2.0) - too_large_penalty - too_small_penalty - edge_penalty


def slot_iou(left: Slot, right: Slot) -> float:
    left_x2 = left.x + left.w
    left_y2 = left.y + left.h
    right_x2 = right.x + right.w
    right_y2 = right.y + right.h

    inter_x1 = max(left.x, right.x)
    inter_y1 = max(left.y, right.y)
    inter_x2 = min(left_x2, right_x2)
    inter_y2 = min(left_y2, right_y2)
    inter_w = max(0, inter_x2 - inter_x1)
    inter_h = max(0, inter_y2 - inter_y1)
    intersection = inter_w * inter_h
    if not intersection:
        return 0.0

    left_area = left.w * left.h
    right_area = right.w * right.h
    return intersection / float(left_area + right_area - intersection)


def slot_is_inside(inner: Slot, outer: Slot) -> bool:
    if inner.w * inner.h >= outer.w * outer.h:
        return False
    inner_cx = inner.x + (inner.w / 2)
    inner_cy = inner.y + (inner.h / 2)
    return outer.x <= inner_cx <= outer.x + outer.w and outer.y <= inner_cy <= outer.y + outer.h


def is_empty_slot(slot_image: Image.Image) -> bool:
    image = np.asarray(slot_image.convert("L"))
    center = image[8:-8, 8:-8] if image.shape[0] > 20 and image.shape[1] > 20 else image
    return float(center.std()) < 13.0


def read_amount(slot_image: Image.Image) -> int:
    width, height = slot_image.size
    count_area = slot_image.crop((int(width * 0.45), int(height * 0.55), width, height))

    gray = cv2.cvtColor(np.asarray(count_area.convert("RGB")), cv2.COLOR_RGB2GRAY)
    enlarged = cv2.resize(gray, None, fx=4, fy=4, interpolation=cv2.INTER_CUBIC)
    _, threshold = cv2.threshold(enlarged, 145, 255, cv2.THRESH_BINARY)

    config = "--psm 7 -c tessedit_char_whitelist=0123456789"
    configure_tesseract()
    try:
        text = pytesseract.image_to_string(threshold, config=config)
    except pytesseract.TesseractNotFoundError as exc:
        raise RuntimeError(tesseract_help_message()) from exc
    digits = re.sub(r"\D+", "", text)

    if not digits:
        return 1

    return max(1, int(digits))


def match_item(slot_image: Image.Image, references: list[ItemReference]) -> tuple[ItemReference, float]:
    best, confidence, _ = match_item_with_candidates(slot_image, references)
    return best, confidence


def match_item_with_candidates(slot_image: Image.Image, references: list[ItemReference]) -> tuple[ItemReference, float, list[str]]:
    scored = sorted(
        ((template_match_score(slot_image, reference), reference) for reference in references),
        key=lambda item: item[0],
        reverse=True,
    )
    best_score, best = scored[0]
    candidates = [reference.name for _, reference in scored[:5]]

    if best_score < 0.45:
        fingerprint = image_fingerprint(crop_slot_icon(slot_image))
        fallback = sorted(
            ((hash_distance(fingerprint, reference), reference) for reference in references),
            key=lambda item: item[0],
        )
        best_distance, best = fallback[0]
        candidates = [reference.name for _, reference in fallback[:5]]
        confidence = max(0.0, min(1.0, 1.0 - (best_distance / 95.0)))
        return best, round(confidence * 0.5, 3), candidates

    return best, round(best_score, 3), candidates


def detect_inventory(
    screenshot_path: Path,
    references: list[ItemReference],
    min_confidence: float,
    read_numbers: bool = True,
) -> list[Detection]:
    screenshot = Image.open(screenshot_path).convert("RGB")
    detections: list[Detection] = []

    for slot in find_slots(screenshot_path):
        # Trim the border so the matcher mostly sees the item.
        pad_x = max(3, int(slot.w * 0.06))
        pad_y = max(3, int(slot.h * 0.06))
        crop = screenshot.crop((slot.x + pad_x, slot.y + pad_y, slot.x + slot.w - pad_x, slot.y + slot.h - pad_y))

        if is_empty_slot(crop):
            continue

        reference, confidence, candidates = match_item_with_candidates(crop, references)
        if confidence < min_confidence:
            continue

        detections.append(
            Detection(
                slot=slot,
                item_name=reference.name,
                amount=read_amount(crop) if read_numbers else 1,
                confidence=confidence,
                image_path=reference.image_path,
                screenshot=screenshot_path,
                candidates=", ".join(candidates),
            )
        )

    return detections


def annotate_screenshot(
    screenshot_path: Path,
    detections: list[Detection],
    output_dir: Path,
) -> Path:
    output_dir.mkdir(parents=True, exist_ok=True)
    image = Image.open(screenshot_path).convert("RGB")
    draw = ImageDraw.Draw(image)

    for detection in detections:
        slot = detection.slot
        label = f"{detection.item_name} x{detection.amount} ({detection.confidence:.2f})"
        draw.rectangle((slot.x, slot.y, slot.x + slot.w, slot.y + slot.h), outline=(60, 220, 90), width=2)
        draw.text((slot.x + 2, max(0, slot.y - 12)), label, fill=(255, 255, 255))

    target = output_dir / f"{screenshot_path.stem}_review.png"
    image.save(target)
    return target


def get_or_create_workbook(path: Path):
    if path.exists():
        return load_workbook(path)

    workbook = Workbook()
    workbook.active.title = "Daily Totals"
    workbook["Daily Totals"].append(["Date", "Item Name", "Amount", "Screenshots", "Average Confidence"])

    raw = workbook.create_sheet("Raw Detections")
    raw.append(["Date", "Screenshot", "Slot X", "Slot Y", "Item Name", "Amount", "Confidence", "Reference PNG"])

    matrix = workbook.create_sheet("By Date")
    matrix.append(["Date"])
    return workbook


def clear_day_rows(sheet, keep_date: str) -> None:
    for row_index in range(sheet.max_row, 1, -1):
        if str(sheet.cell(row=row_index, column=1).value) == keep_date:
            sheet.delete_rows(row_index, 1)


def rebuild_by_date_sheet(workbook) -> None:
    totals = workbook["Daily Totals"]
    matrix = workbook["By Date"]
    matrix.delete_rows(1, matrix.max_row)

    dates: list[str] = []
    items: list[str] = []
    data: dict[tuple[str, str], int] = {}

    for row in totals.iter_rows(min_row=2, values_only=True):
        keep_date, item_name, amount = row[:3]
        if keep_date is None or item_name is None:
            continue
        keep_date = str(keep_date)
        item_name = str(item_name)
        amount = int(amount or 0)
        if keep_date not in dates:
            dates.append(keep_date)
        if item_name not in items:
            items.append(item_name)
        data[(keep_date, item_name)] = data.get((keep_date, item_name), 0) + amount

    items.sort(key=str.casefold)
    matrix.append(["Date", *items])
    for keep_date in sorted(dates):
        matrix.append([keep_date, *[data.get((keep_date, item), 0) for item in items]])

    matrix.freeze_panes = "B2"


def save_to_workbook(
    detections: list[Detection],
    output: Path,
    keep_date: str,
    replace_day: bool,
) -> None:
    workbook = get_or_create_workbook(output)
    totals = workbook["Daily Totals"]
    raw = workbook["Raw Detections"]

    if replace_day:
        clear_day_rows(totals, keep_date)
        clear_day_rows(raw, keep_date)

    grouped: dict[str, dict[str, object]] = {}
    for detection in detections:
        item = grouped.setdefault(
            detection.item_name,
            {"amount": 0, "screenshots": set(), "confidences": []},
        )
        item["amount"] = int(item["amount"]) + detection.amount
        item["screenshots"].add(str(detection.screenshot))  # type: ignore[union-attr]
        item["confidences"].append(detection.confidence)  # type: ignore[union-attr]

        raw.append(
            [
                keep_date,
                str(detection.screenshot),
                detection.slot.x,
                detection.slot.y,
                detection.item_name,
                detection.amount,
                detection.confidence,
                str(detection.image_path),
            ]
        )

    for item_name in sorted(grouped, key=str.casefold):
        item = grouped[item_name]
        confidences = item["confidences"]
        average_confidence = round(sum(confidences) / len(confidences), 3)  # type: ignore[arg-type]
        totals.append(
            [
                keep_date,
                item_name,
                item["amount"],
                ", ".join(sorted(item["screenshots"])),  # type: ignore[arg-type]
                average_confidence,
            ]
        )

    for sheet in (totals, raw):
        sheet.freeze_panes = "A2"
        for column in ("A", "B", "C", "D", "E", "F", "G", "H"):
            sheet.column_dimensions[column].width = 22

    rebuild_by_date_sheet(workbook)
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)


def expand_screenshot_paths(paths: list[Path]) -> list[Path]:
    image_extensions = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
    if not paths and DEFAULT_SCREENSHOTS_DIR.exists():
        paths = [DEFAULT_SCREENSHOTS_DIR]

    expanded: list[Path] = []

    for path in paths:
        if path.is_dir():
            expanded.extend(
                sorted(
                    child for child in path.iterdir()
                    if child.is_file() and child.suffix.lower() in image_extensions
                )
            )
        else:
            expanded.append(path)

    return [path for path in expanded if path.suffix.lower() in image_extensions]


def no_screenshots_message() -> str:
    DEFAULT_SCREENSHOTS_DIR.mkdir(exist_ok=True)
    return (
        "No screenshot files were found.\n\n"
        f"Save your stockpile screenshot as a PNG/JPG in:\n"
        f"  {DEFAULT_SCREENSHOTS_DIR.resolve()}\n\n"
        "Then run:\n"
        "  py New.py\n\n"
        "Or pass a screenshot directly:\n"
        r'  py New.py "C:\Anvil\screenshots\stockpile.png"'
    )


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Read one or more Family Keep inventory screenshots, match the icons "
            "against Anvil_Items PNGs, OCR stack amounts, and append the totals "
            "to an Excel workbook."
        )
    )
    parser.add_argument(
        "screenshots",
        nargs="*",
        type=Path,
        help="Screenshot file(s), or a folder. If omitted, the script scans the screenshots folder.",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        default=DEFAULT_OUTPUT,
        help=f"Output workbook. Default: {DEFAULT_OUTPUT}",
    )
    parser.add_argument(
        "--images-dir",
        type=Path,
        default=DEFAULT_IMAGES_DIR,
        help=f"Item PNG directory made by Anvil.py. Default: {DEFAULT_IMAGES_DIR}",
    )
    parser.add_argument(
        "--items-csv",
        type=Path,
        help="Optional anvil_items.csv to restrict matching to known item names.",
    )
    parser.add_argument(
        "--items-workbook",
        type=Path,
        default=DEFAULT_ITEMS_WORKBOOK,
        help=f"Item workbook made by Anvil.py. Default: {DEFAULT_ITEMS_WORKBOOK}",
    )
    parser.add_argument(
        "--tesseract-cmd",
        help=(
            "Full path to tesseract.exe if it is not on PATH, for example "
            r"C:\Program Files\Tesseract-OCR\tesseract.exe"
        ),
    )
    parser.add_argument(
        "--date",
        default=date.today().isoformat(),
        help="Date to write in the workbook. Default: today.",
    )
    parser.add_argument(
        "--min-confidence",
        type=float,
        default=0.38,
        help="Skip matches below this confidence. Lower keeps more rows; higher is stricter.",
    )
    parser.add_argument(
        "--append",
        action="store_true",
        help="Append to existing rows for the same date instead of replacing that date.",
    )
    parser.add_argument(
        "--annotations-dir",
        type=Path,
        default=DEFAULT_ANNOTATIONS_DIR,
        help="Where review screenshots with match labels are saved.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    configure_tesseract(args.tesseract_cmd)

    references = load_references(args.images_dir, args.items_csv, args.items_workbook)
    all_detections: list[Detection] = []
    screenshots = expand_screenshot_paths(args.screenshots)

    if not screenshots:
        raise RuntimeError(no_screenshots_message())

    print(f"Loaded {len(references)} item references.")
    for screenshot_path in screenshots:
        detections = detect_inventory(screenshot_path, references, args.min_confidence)
        review_path = annotate_screenshot(screenshot_path, detections, args.annotations_dir)
        all_detections.extend(detections)
        print(f"{screenshot_path}: {len(detections)} slots detected. Review: {review_path}")

    if not all_detections:
        raise RuntimeError("No items were detected. Try lowering --min-confidence or check the screenshot crop.")

    save_to_workbook(all_detections, args.output, args.date, replace_day=not args.append)
    print(f"Saved {len(all_detections)} detections to {args.output} for {args.date}.")


if __name__ == "__main__":
    main()
