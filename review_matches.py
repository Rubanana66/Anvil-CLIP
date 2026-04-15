from __future__ import annotations

import argparse
import csv
import re
from datetime import datetime
from pathlib import Path

import cv2
import numpy as np
from openpyxl import Workbook
from openpyxl.drawing.image import Image as ExcelImage
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from PIL import Image, ImageDraw, ImageFont


IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".webp", ".bmp"}
DIGIT_TEMPLATES: dict[str, list[np.ndarray]] | None = None


def item_name_from_icon_path(path: Path) -> str:
    return path.stem.replace("_", " ")


def slot_suffix_from_icon_path(path: Path) -> str:
    # icon_000_r00_c00.png -> 000_r00_c00.png
    return path.name.removeprefix("icon_")


def read_matches(csv_path: Path) -> list[dict[str, str]]:
    with csv_path.open(newline="", encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def load_item_names(icon_dir: Path) -> list[str]:
    names = [
        item_name_from_icon_path(path)
        for path in sorted(icon_dir.iterdir())
        if path.is_file() and path.suffix.lower() in IMAGE_EXTENSIONS
    ]
    return sorted(set(names), key=str.casefold)


def add_image(sheet, image_path: Path, cell: str, width: int, height: int) -> None:
    if not image_path.exists():
        return
    image = ExcelImage(str(image_path))
    image.width = width
    image.height = height
    sheet.add_image(image, cell)


def safe_float(value: str) -> float:
    try:
        return float(value)
    except (TypeError, ValueError):
        return 0.0


def safe_slot_index(path: Path) -> int:
    match = re.search(r"_(\d+)_r\d+_c\d+", path.name)
    return int(match.group(1)) if match else 0


def read_number_crop(number_path: Path) -> str:
    image = cv2.imread(str(number_path), cv2.IMREAD_GRAYSCALE)
    if image is None:
        return ""

    mask = _number_foreground_mask(image)
    image_height, image_width = mask.shape[:2]
    lower_mask = mask.copy()
    lower_mask[: int(image_height * 0.43), :] = 0

    contours, _ = cv2.findContours(lower_mask, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    boxes = []

    for contour in contours:
        x, y, width, height = cv2.boundingRect(contour)
        if width < 3 or height < 8:
            continue
        if height < image_height * 0.15:
            continue
        if y > image_height * 0.92:
            continue
        boxes.append((x, y, width, height))

    boxes = _merge_digit_boxes(boxes)
    if not boxes:
        return ""

    max_bottom = max(y + height for _, y, _, height in boxes)
    boxes = [box for box in boxes if box[1] + box[3] >= max_bottom - 12]

    digits = []
    for box in boxes:
        digit_image = _crop_digit(mask, box)
        digit, score = _match_digit(digit_image)
        if score >= 0.28:
            digits.append(digit)

    return "".join(digits)


def _number_foreground_mask(gray: np.ndarray) -> np.ndarray:
    enlarged = cv2.resize(gray, None, fx=3, fy=3, interpolation=cv2.INTER_CUBIC)
    blurred = cv2.GaussianBlur(enlarged, (3, 3), 0)

    # Stockpile numbers are bright white/gray with dark outline. A high-value
    # mask isolates the digit body well enough for template matching.
    _, mask = cv2.threshold(blurred, 135, 255, cv2.THRESH_BINARY)
    kernel = np.ones((2, 2), dtype=np.uint8)
    mask = cv2.morphologyEx(mask, cv2.MORPH_OPEN, kernel, iterations=1)
    mask = cv2.dilate(mask, kernel, iterations=1)
    return mask


def _merge_digit_boxes(boxes: list[tuple[int, int, int, int]]) -> list[tuple[int, int, int, int]]:
    if not boxes:
        return []

    boxes = sorted(boxes, key=lambda item: item[0])
    merged: list[tuple[int, int, int, int]] = []
    for box in boxes:
        x, y, width, height = box
        if not merged:
            merged.append(box)
            continue

        px, py, pw, ph = merged[-1]
        gap = x - (px + pw)
        vertical_overlap = min(y + height, py + ph) - max(y, py)
        if gap <= 2 and vertical_overlap > min(height, ph) * 0.35:
            x1 = min(px, x)
            y1 = min(py, y)
            x2 = max(px + pw, x + width)
            y2 = max(py + ph, y + height)
            merged[-1] = (x1, y1, x2 - x1, y2 - y1)
        else:
            merged.append(box)

    # Remove tiny leftovers and sort left-to-right.
    median_height = np.median([height for _, _, _, height in merged])
    merged = [box for box in merged if box[3] >= median_height * 0.55]
    return sorted(merged, key=lambda item: item[0])


def _crop_digit(mask: np.ndarray, box: tuple[int, int, int, int]) -> np.ndarray:
    x, y, width, height = box
    pad = 3
    y1 = max(0, y - pad)
    x1 = max(0, x - pad)
    y2 = min(mask.shape[0], y + height + pad)
    x2 = min(mask.shape[1], x + width + pad)
    digit = mask[y1:y2, x1:x2]
    return cv2.resize(digit, (28, 36), interpolation=cv2.INTER_AREA)


def _digit_templates() -> dict[str, list[np.ndarray]]:
    global DIGIT_TEMPLATES
    if DIGIT_TEMPLATES is not None:
        return DIGIT_TEMPLATES

    templates: dict[str, list[np.ndarray]] = {digit: [] for digit in "0123456789"}

    font_paths = [
        Path(r"C:\Windows\Fonts\arial.ttf"),
        Path(r"C:\Windows\Fonts\arialbd.ttf"),
        Path(r"C:\Windows\Fonts\segoeui.ttf"),
        Path(r"C:\Windows\Fonts\segoeuib.ttf"),
        Path(r"C:\Windows\Fonts\tahoma.ttf"),
        Path(r"C:\Windows\Fonts\verdana.ttf"),
    ]
    for font_path in font_paths:
        if not font_path.exists():
            continue
        for font_size in range(24, 42, 2):
            try:
                font = ImageFont.truetype(str(font_path), font_size)
            except OSError:
                continue
            for stroke_width in (0, 1):
                for digit in "0123456789":
                    canvas = Image.new("L", (60, 70), 0)
                    draw = ImageDraw.Draw(canvas)
                    draw.text(
                        (8, 4),
                        digit,
                        font=font,
                        fill=255,
                        stroke_width=stroke_width,
                        stroke_fill=255,
                    )
                    template = _normalise_template(np.asarray(canvas))
                    if template is not None:
                        templates[digit].append(template)

    # Fallback templates for non-Windows machines.
    for digit in "0123456789":
        canvas = np.zeros((48, 36), dtype=np.uint8)
        cv2.putText(
            canvas,
            digit,
            (3, 39),
            cv2.FONT_HERSHEY_SIMPLEX,
            1.25,
            255,
            2,
            cv2.LINE_AA,
        )
        template = _normalise_template(canvas)
        if template is not None:
            templates[digit].append(template)

    DIGIT_TEMPLATES = templates
    return templates


def _normalise_template(template: np.ndarray) -> np.ndarray | None:
    _, binary = cv2.threshold(template, 30, 255, cv2.THRESH_BINARY)
    y, x = np.where(binary > 0)
    if len(x) == 0 or len(y) == 0:
        return None
    cropped = binary[y.min() : y.max() + 1, x.min() : x.max() + 1]
    return cv2.resize(cropped, (28, 36), interpolation=cv2.INTER_AREA)


def _match_digit(digit_image: np.ndarray) -> tuple[str, float]:
    best_digit = ""
    best_score = -1.0
    query = digit_image.astype(np.float32) / 255.0

    for digit, templates in _digit_templates().items():
        for template in templates:
            target = template.astype(np.float32) / 255.0
            score = float(cv2.matchTemplate(query, target, cv2.TM_CCOEFF_NORMED)[0][0])
            if score > best_score:
                best_digit = digit
                best_score = score

    return best_digit, best_score


def build_review_workbook(
    parser_output_dir: Path,
    itemlist_dir: Path = Path("itemlist"),
    output_path: Path | None = None,
) -> Path:
    matches_csv = parser_output_dir / "icon_matches" / "icon_matches.csv"
    if not matches_csv.exists():
        raise FileNotFoundError(f"Missing icon matcher CSV: {matches_csv}")

    output_path = output_path or parser_output_dir / "review_icon_matches.xlsx"
    matches = read_matches(matches_csv)
    item_names = load_item_names(itemlist_dir)

    workbook = Workbook()
    review = workbook.active
    review.title = "Review"

    headers = [
        "Slot",
        "Slot Crop",
        "Icon Crop",
        "Detected Amount",
        "Best Item",
        "Similarity",
        "Review Needed",
        "Top Candidates",
        "Correct Item",
        "Correct Amount",
        "Notes",
        "Icon Path",
        "Slot Path",
        "Number Path",
    ]
    review.append(headers)

    header_fill = PatternFill("solid", fgColor="1F2937")
    header_font = Font(color="FFFFFF", bold=True)
    for cell in review[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row_index, match in enumerate(matches, start=2):
        icon_path = Path(match["query_path"])
        if not icon_path.is_absolute():
            icon_path = Path.cwd() / icon_path

        suffix = slot_suffix_from_icon_path(icon_path)
        slot_path = parser_output_dir / "slots" / f"slot_{suffix}"
        number_path = parser_output_dir / "numbers" / f"number_{suffix}"
        detected_amount = read_number_crop(number_path)
        slot_index = safe_slot_index(icon_path)
        best_item = match.get("best_item", "")
        similarity = safe_float(match.get("similarity", ""))
        review_needed = str(match.get("review_needed", "")).lower() == "true"

        review.append(
            [
                slot_index,
                "",
                "",
                detected_amount,
                best_item,
                similarity,
                review_needed,
                match.get("top_candidates", ""),
                best_item,
                detected_amount,
                "",
                str(icon_path),
                str(slot_path),
                str(number_path),
            ]
        )
        review.row_dimensions[row_index].height = 74

        add_image(review, slot_path, f"B{row_index}", width=72, height=72)
        add_image(review, icon_path, f"C{row_index}", width=72, height=72)

        review.cell(row=row_index, column=6).number_format = "0.0000"
        for col in range(1, len(headers) + 1):
            review.cell(row=row_index, column=col).alignment = Alignment(vertical="top", wrap_text=True)

    items_sheet = workbook.create_sheet("Item Names")
    for row_index, item_name in enumerate(item_names, start=1):
        items_sheet.cell(row=row_index, column=1, value=item_name)
    items_sheet.sheet_state = "hidden"

    if item_names:
        validation = DataValidation(
            type="list",
            formula1=f"'Item Names'!$A$1:$A${len(item_names)}",
            allow_blank=True,
        )
        review.add_data_validation(validation)
        validation.add(f"I2:I{max(2, len(matches) + 1)}")

    widths = {
        "A": 8,
        "B": 14,
        "C": 14,
        "D": 14,
        "E": 28,
        "F": 12,
        "G": 14,
        "H": 72,
        "I": 28,
        "J": 14,
        "K": 32,
        "L": 50,
        "M": 50,
        "N": 50,
    }
    for column, width in widths.items():
        review.column_dimensions[column].width = width

    review.freeze_panes = "A2"
    review.auto_filter.ref = review.dimensions
    output_path.parent.mkdir(parents=True, exist_ok=True)
    try:
        workbook.save(output_path)
    except PermissionError:
        fallback = output_path.with_name(
            f"{output_path.stem}_{datetime.now().strftime('%Y%m%d_%H%M%S')}{output_path.suffix}"
        )
        workbook.save(fallback)
        return fallback
    return output_path


def latest_parser_output(base_dir: Path = Path("output_debug")) -> Path:
    candidates = [
        path for path in base_dir.iterdir()
        if path.is_dir() and (path / "icon_matches" / "icon_matches.csv").exists()
    ]
    if not candidates:
        raise FileNotFoundError(f"No parser output with icon_matches.csv found under {base_dir}")
    return max(candidates, key=lambda path: path.stat().st_mtime)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Create an Excel review workbook for OpenCLIP icon matches.")
    parser.add_argument(
        "parser_output_dir",
        nargs="?",
        type=Path,
        help="Parser output directory, for example output_debug/Townkeep_20260413_181451. Defaults to latest.",
    )
    parser.add_argument(
        "--itemlist-dir",
        type=Path,
        default=Path("itemlist"),
        help="Folder containing clean item PNGs. Default: itemlist",
    )
    parser.add_argument(
        "-o",
        "--output",
        type=Path,
        help="Output .xlsx path. Default: review_icon_matches.xlsx inside parser output dir.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    parser_output_dir = args.parser_output_dir or latest_parser_output()
    output = build_review_workbook(parser_output_dir, args.itemlist_dir, args.output)
    print(f"Review workbook written to {output}")


if __name__ == "__main__":
    main()
