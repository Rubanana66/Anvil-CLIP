from __future__ import annotations

import argparse
import csv
import re
import shutil
import sys
import time
from io import BytesIO
from pathlib import Path
from typing import Iterable

try:
    import requests
    from bs4 import BeautifulSoup
    from PIL import Image
except ImportError as exc:
    missing = exc.name or "a required package"
    print(
        f"Missing dependency: {missing}\n\n"
        "Install dependencies with:\n"
        "  py -m pip install -r requirements.txt\n\n"
        "Then run:\n"
        "  py download_anvil_items.py\n",
        file=sys.stderr,
    )
    raise SystemExit(1)


WIKI_API = "https://anvilempires.wiki.gg/api.php"
WIKI_BASE = "https://anvilempires.wiki.gg"

DEFAULT_CATEGORIES = (
    "Category:Armours",
    "Category:Clothing",
    "Category:Consumables",
    "Category:Equipment",
    "Category:Foods",
    "Category:Items",
    "Category:Large Items",
    "Category:Materials",
    "Category:Resources",
    "Category:Tools",
    "Category:Weapons",
)

EXCLUDED_CATEGORY_WORDS = {
    "building",
    "buildings",
    "construction",
    "crafting stations",
    "decorations",
    "doors",
    "fortifications",
    "foundations",
    "furniture",
    "gates",
    "production buildings",
    "ships",
    "siege",
    "structures",
    "transport",
    "vehicles",
    "walls",
}

# Historical list from the first downloader version. It is intentionally not
# used for item filtering anymore because it skipped valid items such as
# "Mill Harness" and "Auxiliary Chest Armour".
EXCLUDED_TITLE_WORDS = {
    "bank",
    "barracks",
    "bench",
    "building",
    "campfire",
    "cart",
    "castle",
    "chest",
    "door",
    "fence",
    "forge",
    "furnace",
    "gate",
    "hall",
    "house",
    "kiln",
    "mill",
    "mine",
    "oven",
    "palisade",
    "ship",
    "shop",
    "stable",
    "station",
    "structure",
    "tower",
    "wall",
    "warehouse",
    "workbench",
    "workshop",
}

EXCLUDED_EXACT_TITLES = {
    "Armours",
    "Foods",
    "Inventory",
    "Items",
    "Mill",
    "Resources",
    "Tools",
    "Weapons",
}

UNRELEASED_WARNING_PHRASES = (
    "unreleased or disabled content",
    "doesn't apply to the current live version",
    "doesn’t apply to the current live version",
)

IMAGE_SKIP_WORDS = {
    "ambox",
    "arrow",
    "background",
    "blank",
    "edit",
    "external",
    "iconbackground",
    "logo",
    "placeholder",
    "search",
    "sprite",
}


def api_get(session: requests.Session, **params: object) -> dict:
    params.setdefault("format", "json")
    params.setdefault("formatversion", "2")

    for attempt in range(5):
        response = session.get(WIKI_API, params=params, timeout=30)
        if response.status_code == 429:
            time.sleep(2 + attempt)
            continue
        response.raise_for_status()
        data = response.json()
        if data.get("error", {}).get("code") == "maxlag":
            time.sleep(2 + attempt)
            continue
        return data

    raise RuntimeError(f"Wiki API did not answer cleanly after retries: {params}")


def iter_category_members(
    session: requests.Session,
    category: str,
    seen_categories: set[str],
) -> Iterable[str]:
    if category in seen_categories:
        return
    seen_categories.add(category)

    query = {
        "action": "query",
        "list": "categorymembers",
        "cmtitle": category,
        "cmlimit": "500",
        "cmtype": "page|subcat",
    }

    while True:
        data = api_get(session, **query)
        for member in data.get("query", {}).get("categorymembers", []):
            title = member["title"]
            if member.get("ns") == 14:
                yield from iter_category_members(session, title, seen_categories)
            elif member.get("ns") == 0:
                yield title

        continuation = data.get("continue")
        if not continuation:
            break
        query.update(continuation)


def candidate_titles(session: requests.Session, categories: Iterable[str]) -> list[str]:
    titles: set[str] = set()
    seen_categories: set[str] = set()

    for category in categories:
        print(f"Scanning {category}...")
        titles.update(iter_category_members(session, category, seen_categories))

    return sorted(titles, key=str.casefold)


def parse_page(session: requests.Session, title: str) -> tuple[str, list[str]]:
    data = api_get(
        session,
        action="parse",
        page=title,
        prop="text|categories",
        redirects="1",
        disablelimitreport="1",
    )
    parsed = data.get("parse", {})
    html = parsed.get("text", "")
    categories = [cat.get("category", "") for cat in parsed.get("categories", [])]
    return html, categories


def is_building_or_overview_page(title: str, categories: Iterable[str], html: str = "") -> tuple[bool, str]:
    if title in EXCLUDED_EXACT_TITLES:
        return True, "overview_or_disambiguation"

    folded_categories = " ".join(categories).lower()
    if any(word in folded_categories for word in EXCLUDED_CATEGORY_WORDS):
        return True, "building_or_structure_category"

    if is_unreleased_or_disabled(html, categories):
        return True, "unreleased_or_disabled"

    return False, ""


def is_unreleased_or_disabled(html: str, categories: Iterable[str]) -> bool:
    text = BeautifulSoup(html, "html.parser").get_text(" ", strip=True).lower()
    folded_categories = " ".join(categories).lower()
    haystack = f"{text} {folded_categories}"
    return any(phrase in haystack for phrase in UNRELEASED_WARNING_PHRASES)


def absolute_image_url(src: str | None) -> str | None:
    if not src:
        return None
    if src.startswith("//"):
        return normalize_wiki_image_url("https:" + src)
    if src.startswith("/"):
        return normalize_wiki_image_url(WIKI_BASE + src)
    if src.startswith("http://") or src.startswith("https://"):
        return normalize_wiki_image_url(src)
    return None


def normalize_wiki_image_url(url: str) -> str:
    base, separator, query = url.partition("?")
    match = re.match(r"^(?P<prefix>https?://[^/]+/images)/thumb/(?P<file_path>.+?)/[^/]+$", base)
    if not match:
        return url

    original = f"{match.group('prefix')}/{match.group('file_path')}"
    return f"{original}{separator}{query}" if separator else original


def image_score(img) -> int:
    src = img.get("src") or img.get("data-src") or ""
    alt = img.get("alt", "")
    text = f"{src} {alt}".lower()

    if not src:
        return -100
    if any(skip in text for skip in IMAGE_SKIP_WORDS):
        return -50

    score = 0
    if ".png" in text:
        score += 20
    if "images" in text or "thumb" in text:
        score += 10
    if img.find_parent(class_=re.compile("infobox|portable-infobox|pi-image", re.I)):
        score += 30

    try:
        width = int(img.get("width") or 0)
        height = int(img.get("height") or 0)
    except ValueError:
        width = height = 0

    if width >= 24 and height >= 24:
        score += 10
    if width > 350 or height > 350:
        score -= 15

    return score


def find_item_image_url(html: str) -> str | None:
    soup = BeautifulSoup(html, "html.parser")

    candidates = soup.select(".infobox img, .portable-infobox img, .pi-image img")
    if not candidates:
        candidates = soup.select(".mw-parser-output img")

    if not candidates:
        return None

    best = max(candidates, key=image_score)
    if image_score(best) < 0:
        return None

    return absolute_image_url(best.get("src") or best.get("data-src"))


def safe_filename(item_name: str) -> str:
    slug = re.sub(r"[^A-Za-z0-9._-]+", "_", item_name).strip("_")
    return f"{slug}.png"


def download_png(session: requests.Session, image_url: str, output_path: Path) -> None:
    response = session.get(image_url, timeout=30)
    response.raise_for_status()

    with Image.open(BytesIO(response.content)) as image:
        output_path.parent.mkdir(parents=True, exist_ok=True)
        image.convert("RGBA").save(output_path, "PNG")


def write_manifest(rows: list[dict[str, str]], output_dir: Path) -> None:
    manifest = output_dir / "itemlist_manifest.csv"
    with manifest.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.DictWriter(
            handle,
            fieldnames=["Item Name", "Status", "Reason", "PNG File", "Wiki Page", "Image URL"],
        )
        writer.writeheader()
        writer.writerows(rows)
    print(f"Manifest written to {manifest}")


def titles_from_xlsx(path: Path, name_column: str = "Name") -> list[str]:
    try:
        from openpyxl import load_workbook
    except ImportError as exc:
        print(
            "Missing dependency: openpyxl\n\n"
            "Install dependencies with:\n"
            "  py -m pip install -r requirements.txt\n",
            file=sys.stderr,
        )
        raise SystemExit(1) from exc

    if not path.exists():
        raise FileNotFoundError(f"Item workbook not found: {path}")

    workbook = load_workbook(path, data_only=True, read_only=True)
    wanted = name_column.strip().lower()

    for worksheet in workbook.worksheets:
        max_header_row = min(20, worksheet.max_row)
        for row_index in range(1, max_header_row + 1):
            headers = [
                str(worksheet.cell(row_index, column).value or "").strip()
                for column in range(1, worksheet.max_column + 1)
            ]
            lowered = [header.lower() for header in headers]
            if wanted not in lowered:
                continue

            name_column_index = lowered.index(wanted) + 1
            names: list[str] = []
            seen: set[str] = set()
            for data_row in range(row_index + 1, worksheet.max_row + 1):
                raw_name = worksheet.cell(data_row, name_column_index).value
                if raw_name is None:
                    continue
                item_name = str(raw_name).strip()
                if not item_name or item_name.lower() in seen:
                    continue
                seen.add(item_name.lower())
                names.append(item_name)

            if names:
                print(f"Loaded {len(names)} item names from {path} sheet '{worksheet.title}'.")
                return names

    raise ValueError(f"Could not find a '{name_column}' column in {path}")


def archive_pngs_not_in_source(output_dir: Path, titles: Iterable[str]) -> int:
    allowed_names = {safe_filename(title).lower() for title in titles}
    stale_paths = [
        path
        for path in output_dir.glob("*.png")
        if path.name.lower() not in allowed_names
    ]
    if not stale_paths:
        return 0

    archive_dir = output_dir / f"_archived_not_in_source_{datetime_stamp()}"
    archive_dir.mkdir(parents=True, exist_ok=True)
    for path in stale_paths:
        shutil.move(str(path), str(archive_dir / path.name))
    print(f"Archived {len(stale_paths)} PNG(s) not present in the source item list to {archive_dir}")
    return len(stale_paths)


def archive_existing_png(output_dir: Path, title: str, reason: str) -> None:
    png_path = output_dir / safe_filename(title)
    if not png_path.exists():
        return

    archive_dir = output_dir / f"_archived_{reason}_{datetime_stamp()}"
    archive_dir.mkdir(parents=True, exist_ok=True)
    shutil.move(str(png_path), str(archive_dir / png_path.name))
    print(f"  Archived {png_path.name} because {reason}.")


def datetime_stamp() -> str:
    return time.strftime("%Y%m%d_%H%M%S")


def download_all_items(
    output_dir: Path,
    categories: Iterable[str],
    limit: int | None = None,
    source_xlsx: Path | None = None,
    name_column: str = "Name",
    force: bool = False,
    archive_not_in_source: bool = False,
) -> None:
    session = requests.Session()
    session.headers.update(
        {
            "User-Agent": (
                "AnvilItemPngDownloader/1.0 "
                "(local script; source https://anvilempires.wiki.gg/)"
            )
        }
    )

    if source_xlsx:
        titles = titles_from_xlsx(source_xlsx, name_column)
    else:
        titles = candidate_titles(session, categories)

    if limit:
        titles = titles[:limit]

    rows: list[dict[str, str]] = []
    skipped = 0
    downloaded_or_found = 0

    output_dir.mkdir(parents=True, exist_ok=True)

    if source_xlsx and archive_not_in_source:
        archive_pngs_not_in_source(output_dir, titles)

    for index, title in enumerate(titles, start=1):
        print(f"[{index}/{len(titles)}] {title}")
        try:
            if title in EXCLUDED_EXACT_TITLES:
                rows.append(
                    {
                        "Item Name": title,
                        "Status": "skipped",
                        "Reason": "overview_or_disambiguation",
                        "PNG File": "",
                        "Wiki Page": f"{WIKI_BASE}/wiki/{title.replace(' ', '_')}",
                        "Image URL": "",
                    }
                )
                skipped += 1
                continue

            html, page_categories = parse_page(session, title)
            should_skip, reason = is_building_or_overview_page(title, page_categories, html)
            if should_skip:
                if archive_not_in_source:
                    archive_existing_png(output_dir, title, reason)
                rows.append(
                    {
                        "Item Name": title,
                        "Status": "skipped",
                        "Reason": reason,
                        "PNG File": "",
                        "Wiki Page": f"{WIKI_BASE}/wiki/{title.replace(' ', '_')}",
                        "Image URL": "",
                    }
                )
                skipped += 1
                continue

            image_url = find_item_image_url(html)
            if not image_url:
                if archive_not_in_source:
                    archive_existing_png(output_dir, title, "missing_image")
                rows.append(
                    {
                        "Item Name": title,
                        "Status": "missing_image",
                        "Reason": "no_usable_png_found",
                        "PNG File": "",
                        "Wiki Page": f"{WIKI_BASE}/wiki/{title.replace(' ', '_')}",
                        "Image URL": "",
                    }
                )
                skipped += 1
                continue

            png_path = output_dir / safe_filename(title)
            status = "found"
            if force or not png_path.exists() or png_path.stat().st_size == 0:
                download_png(session, image_url, png_path)
                status = "downloaded"
            downloaded_or_found += 1

            rows.append(
                {
                    "Item Name": title,
                    "Status": status,
                    "Reason": "",
                    "PNG File": str(png_path),
                    "Wiki Page": f"{WIKI_BASE}/wiki/{title.replace(' ', '_')}",
                    "Image URL": image_url,
                }
            )
        except Exception as exc:
            print(f"  Skipped {title}: {exc}", file=sys.stderr)
            existing_png = output_dir / safe_filename(title)
            status = "found_local_after_error" if existing_png.exists() else "error"
            rows.append(
                {
                    "Item Name": title,
                    "Status": status,
                    "Reason": str(exc),
                    "PNG File": str(existing_png) if existing_png.exists() else "",
                    "Wiki Page": f"{WIKI_BASE}/wiki/{title.replace(' ', '_')}",
                    "Image URL": "",
                }
            )
            skipped += 1

    write_manifest(rows, output_dir)
    print(f"Done. Downloaded/found {downloaded_or_found} item PNGs. Skipped/reported {skipped} pages.")


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description="Download Anvil Empires item PNGs from wiki.gg into an itemlist folder."
    )
    parser.add_argument(
        "-o",
        "--output-dir",
        type=Path,
        default=Path("itemlist"),
        help="Folder where item PNGs are saved. Default: itemlist",
    )
    parser.add_argument(
        "--category",
        action="append",
        default=[],
        help="Extra category to scan, for example 'Category:Resources'. Can be used more than once.",
    )
    parser.add_argument(
        "--only-categories",
        action="store_true",
        help="Only scan categories passed with --category instead of the default item categories.",
    )
    parser.add_argument(
        "--limit",
        type=int,
        help="Process only the first N candidate pages. Useful for testing.",
    )
    parser.add_argument(
        "--source-xlsx",
        type=Path,
        help="Use item names from an Excel workbook instead of scanning category pages.",
    )
    parser.add_argument(
        "--name-column",
        default="Name",
        help="Column name to read when --source-xlsx is used. Default: Name",
    )
    parser.add_argument(
        "--force",
        action="store_true",
        help="Download PNGs again even if the target file already exists.",
    )
    parser.add_argument(
        "--archive-not-in-source",
        action="store_true",
        help="Move PNGs whose filenames are not in --source-xlsx to an archive folder inside the output dir.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    categories = args.category if args.only_categories else [*DEFAULT_CATEGORIES, *args.category]
    download_all_items(
        output_dir=args.output_dir,
        categories=categories,
        limit=args.limit,
        source_xlsx=args.source_xlsx,
        name_column=args.name_column,
        force=args.force,
        archive_not_in_source=args.archive_not_in_source,
    )


if __name__ == "__main__":
    main()
