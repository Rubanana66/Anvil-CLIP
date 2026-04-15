from __future__ import annotations

import argparse
import csv
import hashlib
import logging
import re
from dataclasses import dataclass
from datetime import datetime
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font, PatternFill


LOGGER = logging.getLogger("anvil_finalize_review")


@dataclass(frozen=True)
class CleanRow:
    screenshot: str
    slot: int
    grid_row: int
    grid_col: int
    item: str
    amount: int | None
    detected_item: str
    detected_amount: str
    icon_confidence: float | None
    review_needed: bool
    notes: str
    parser_output: str
    icon_path: str
    slot_path: str
    number_path: str

    def as_list(self) -> list[Any]:
        return [
            self.screenshot,
            self.slot,
            self.grid_row,
            self.grid_col,
            self.item,
            self.amount,
            self.detected_item,
            self.detected_amount,
            self.icon_confidence,
            self.review_needed,
            self.notes,
            self.parser_output,
            self.icon_path,
            self.slot_path,
            self.number_path,
        ]


def latest_review_workbook(output_dir: Path = Path("output_debug")) -> Path:
    candidates = sorted(
        output_dir.glob("combined_review_*/combined_stockpile_review.xlsx"),
        key=lambda path: path.stat().st_mtime,
        reverse=True,
    )
    if not candidates:
        raise FileNotFoundError(f"No combined review workbook found under {output_dir}")
    return candidates[0]


def read_clean_rows(workbook_path: Path) -> list[CleanRow]:
    workbook = load_workbook(workbook_path, data_only=True)
    if "Review" not in workbook.sheetnames:
        raise ValueError(f"Workbook has no Review sheet: {workbook_path}")

    sheet = workbook["Review"]
    headers = {
        str(cell.value).strip(): index
        for index, cell in enumerate(sheet[1], start=1)
        if cell.value is not None
    }

    required = ["Screenshot", "Slot", "Correct Item", "Correct Amount"]
    missing = [header for header in required if header not in headers]
    if missing:
        raise ValueError(f"Missing required columns in Review sheet: {', '.join(missing)}")

    rows: list[CleanRow] = []
    for row_number in range(2, sheet.max_row + 1):
        screenshot = _cell(sheet, row_number, headers, "Screenshot")
        if not screenshot:
            continue

        detected_item = _cell(sheet, row_number, headers, "Detected Item")
        detected_amount = _cell(sheet, row_number, headers, "Detected Amount")
        correct_item = _cell(sheet, row_number, headers, "Correct Item") or detected_item
        correct_amount_raw = _cell(sheet, row_number, headers, "Correct Amount") or detected_amount

        item = str(correct_item).strip()
        amount = parse_amount(correct_amount_raw)
        review_needed = parse_bool(_cell(sheet, row_number, headers, "Review Needed"))

        if not item or amount is None:
            review_needed = True

        rows.append(
            CleanRow(
                screenshot=str(screenshot).strip(),
                slot=parse_int(_cell(sheet, row_number, headers, "Slot")) or 0,
                grid_row=parse_int(_cell(sheet, row_number, headers, "Row")) or 0,
                grid_col=parse_int(_cell(sheet, row_number, headers, "Col")) or 0,
                item=item,
                amount=amount,
                detected_item=str(detected_item or "").strip(),
                detected_amount=str(detected_amount or "").strip(),
                icon_confidence=parse_float(_cell(sheet, row_number, headers, "Icon Confidence")),
                review_needed=review_needed,
                notes=str(_cell(sheet, row_number, headers, "Notes") or "").strip(),
                parser_output=str(_cell(sheet, row_number, headers, "Parser Output") or "").strip(),
                icon_path=str(_cell(sheet, row_number, headers, "Icon Path") or "").strip(),
                slot_path=str(_cell(sheet, row_number, headers, "Slot Path") or "").strip(),
                number_path=str(_cell(sheet, row_number, headers, "Number Path") or "").strip(),
            )
        )

    return rows


def build_totals(rows: list[CleanRow]) -> list[tuple[str, int, int]]:
    totals: dict[str, int] = {}
    row_counts: dict[str, int] = {}

    for row in rows:
        if not row.item or row.amount is None:
            continue
        totals[row.item] = totals.get(row.item, 0) + row.amount
        row_counts[row.item] = row_counts.get(row.item, 0) + 1

    return sorted(
        ((item, amount, row_counts[item]) for item, amount in totals.items()),
        key=lambda item: item[0].casefold(),
    )


def find_duplicate_screenshots(rows: list[CleanRow], screenshots_dir: Path) -> list[tuple[str, list[str]]]:
    screenshots = sorted({row.screenshot for row in rows})
    by_hash: dict[str, list[str]] = {}

    for screenshot in screenshots:
        path = screenshots_dir / screenshot
        if not path.exists() or not path.is_file():
            continue
        digest = file_sha256(path)
        by_hash.setdefault(digest, []).append(screenshot)

    return [
        (digest, names)
        for digest, names in sorted(by_hash.items(), key=lambda item: item[1][0].casefold())
        if len(names) > 1
    ]


def write_outputs(
    review_workbook: Path,
    rows: list[CleanRow],
    screenshots_dir: Path,
    output_dir: Path | None,
) -> tuple[Path, Path, Path]:
    output_dir = output_dir or review_workbook.parent / "final"
    output_dir.mkdir(parents=True, exist_ok=True)

    clean_xlsx = output_dir / "final_stockpile_inventory.xlsx"
    clean_csv = output_dir / "final_stockpile_rows.csv"
    totals_csv = output_dir / "item_totals.csv"

    write_clean_csv(clean_csv, rows)
    write_totals_csv(totals_csv, build_totals(rows))
    write_clean_workbook(clean_xlsx, review_workbook, rows, screenshots_dir)
    return clean_xlsx, clean_csv, totals_csv


def write_clean_workbook(
    output_path: Path,
    source_workbook: Path,
    rows: list[CleanRow],
    screenshots_dir: Path,
) -> None:
    workbook = Workbook()
    reviewed = workbook.active
    reviewed.title = "Reviewed Rows"

    reviewed_headers = [
        "Screenshot",
        "Slot",
        "Row",
        "Col",
        "Item",
        "Amount",
        "Detected Item",
        "Detected Amount",
        "Icon Confidence",
        "Review Needed",
        "Notes",
        "Parser Output",
        "Icon Path",
        "Slot Path",
        "Number Path",
    ]
    reviewed.append(reviewed_headers)
    _style_header(reviewed)
    for row in rows:
        reviewed.append(row.as_list())

    totals = workbook.create_sheet("Item Totals")
    totals.append(["Item", "Total Amount", "Rows"])
    _style_header(totals)
    for item, amount, row_count in build_totals(rows):
        totals.append([item, amount, row_count])

    metadata = workbook.create_sheet("Metadata")
    metadata.append(["Key", "Value"])
    _style_header(metadata)
    metadata.append(["Created At", datetime.now().isoformat(timespec="seconds")])
    metadata.append(["Source Workbook", str(source_workbook)])
    metadata.append(["Reviewed Rows", len(rows)])
    metadata.append(["Unique Items", len(build_totals(rows))])
    metadata.append(["Rows Still Needing Review", sum(1 for row in rows if row.review_needed)])

    duplicates = find_duplicate_screenshots(rows, screenshots_dir)
    duplicate_sheet = workbook.create_sheet("Duplicate Screenshots")
    duplicate_sheet.append(["SHA256", "Screenshot Names"])
    _style_header(duplicate_sheet)
    for digest, names in duplicates:
        duplicate_sheet.append([digest, " | ".join(names)])

    _format_sheet(reviewed, widths={
        "A": 42,
        "B": 8,
        "C": 8,
        "D": 8,
        "E": 30,
        "F": 12,
        "G": 30,
        "H": 16,
        "I": 16,
        "J": 15,
        "K": 36,
        "L": 58,
        "M": 58,
        "N": 58,
        "O": 58,
    })
    _format_sheet(totals, widths={"A": 34, "B": 16, "C": 12})
    _format_sheet(metadata, widths={"A": 24, "B": 90})
    _format_sheet(duplicate_sheet, widths={"A": 72, "B": 100})

    workbook.save(output_path)


def write_clean_csv(output_path: Path, rows: list[CleanRow]) -> None:
    headers = [
        "screenshot",
        "slot",
        "row",
        "col",
        "item",
        "amount",
        "detected_item",
        "detected_amount",
        "icon_confidence",
        "review_needed",
        "notes",
        "parser_output",
        "icon_path",
        "slot_path",
        "number_path",
    ]
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(headers)
        for row in rows:
            writer.writerow(row.as_list())


def write_totals_csv(output_path: Path, totals: list[tuple[str, int, int]]) -> None:
    with output_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(["item", "total_amount", "rows"])
        writer.writerows(totals)


def parse_amount(value: Any) -> int | None:
    if value is None:
        return None
    if isinstance(value, bool):
        return None
    if isinstance(value, int):
        return value
    if isinstance(value, float):
        return int(value) if value.is_integer() else round(value)

    text = str(value).strip()
    if not text:
        return None
    text = text.replace(",", "")
    match = re.search(r"-?\d+", text)
    if not match:
        return None
    return int(match.group(0))


def parse_int(value: Any) -> int | None:
    amount = parse_amount(value)
    return amount


def parse_float(value: Any) -> float | None:
    if value is None or value == "":
        return None
    try:
        return float(value)
    except (TypeError, ValueError):
        return None


def parse_bool(value: Any) -> bool:
    if isinstance(value, bool):
        return value
    if value is None:
        return False
    return str(value).strip().lower() in {"true", "yes", "1", "y"}


def file_sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _cell(sheet: Any, row_number: int, headers: dict[str, int], header: str) -> Any:
    column = headers.get(header)
    if column is None:
        return None
    return sheet.cell(row=row_number, column=column).value


def _style_header(sheet: Any) -> None:
    fill = PatternFill("solid", fgColor="1F2937")
    font = Font(color="FFFFFF", bold=True)
    for cell in sheet[1]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)


def _format_sheet(sheet: Any, widths: dict[str, int]) -> None:
    for column, width in widths.items():
        sheet.column_dimensions[column].width = width
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for row in sheet.iter_rows():
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(description="Create clean inventory exports from a corrected review workbook.")
    parser.add_argument(
        "review_workbook",
        nargs="?",
        type=Path,
        help="Path to combined_stockpile_review.xlsx. Defaults to the newest combined review workbook.",
    )
    parser.add_argument(
        "--screenshots-dir",
        type=Path,
        default=Path("screenshots"),
        help="Folder containing the original screenshots. Used only to detect duplicate screenshot files.",
    )
    parser.add_argument(
        "--output-dir",
        type=Path,
        help="Output folder. Default: final/ inside the review workbook folder.",
    )
    parser.add_argument(
        "--log-level",
        default="INFO",
        choices=("DEBUG", "INFO", "WARNING", "ERROR"),
        help="Logging level. Default: INFO",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    logging.basicConfig(level=getattr(logging, args.log_level), format="%(levelname)s: %(message)s")

    review_workbook = args.review_workbook or latest_review_workbook()
    LOGGER.info("Using review workbook: %s", review_workbook)

    rows = read_clean_rows(review_workbook)
    clean_xlsx, clean_csv, totals_csv = write_outputs(
        review_workbook=review_workbook,
        rows=rows,
        screenshots_dir=args.screenshots_dir,
        output_dir=args.output_dir,
    )

    rows_needing_review = sum(1 for row in rows if row.review_needed)
    totals = build_totals(rows)

    print()
    print("Anvil finalized stockpile export")
    print("--------------------------------")
    print(f"Source review workbook: {review_workbook}")
    print(f"Reviewed rows: {len(rows)}")
    print(f"Unique items: {len(totals)}")
    print(f"Rows still marked review: {rows_needing_review}")
    print(f"Excel: {clean_xlsx}")
    print(f"Rows CSV: {clean_csv}")
    print(f"Totals CSV: {totals_csv}")


if __name__ == "__main__":
    main()
